"""Arma el respaldo diario de "Mis Finanzas del Mes" como planilla Excel.

Uso: python3 respaldo.py <carpeta> <salida.xlsx>
<carpeta> tiene las subcarpetas pagos/, trabajos/ y gastos/ con un JSON por
documento (lo que deja ArtifactData list con out_dir).
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

IVA = 0.19
COLORES = {"Pagos": "8FB3E0", "Trabajos": "8CC9A8", "Gastos": "E8B98C"}


def leer(carpeta, coleccion):
    filas = []
    for f in sorted((Path(carpeta) / coleccion).glob("*.json")):
        d = json.loads(f.read_text(encoding="utf-8"))
        d["_id"] = f.stem
        filas.append(d)
    return filas


def neto(t):
    if t.get("tipo") == "trato":
        return t.get("valor", 0)
    return t.get("cantidad", 0) * t.get("valor", 0)


def hoja(wb, nombre, cabecera, filas, anchos):
    ws = wb.create_sheet(nombre)
    ws.append(cabecera)
    for celda in ws[1]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill("solid", fgColor=COLORES[nombre])
    for fila in filas:
        ws.append(fila)
    for i, ancho in enumerate(anchos):
        ws.column_dimensions[chr(65 + i)].width = ancho
    ws.freeze_panes = "A2"


def main(carpeta, salida):
    pagos = leer(carpeta, "pagos")
    trabajos = sorted(leer(carpeta, "trabajos"), key=lambda t: t.get("fecha", ""))
    gastos = sorted(leer(carpeta, "gastos"), key=lambda g: g.get("fecha", ""))

    wb = Workbook()
    wb.remove(wb.active)
    hoja(wb, "Pagos",
         ["Pago", "Monto", "Día", "Cuotas en total", "Primera cuota (mes)", "Desde (mes)", "Meses pagados", "Ejemplo"],
         [[p.get("nombre"), p.get("monto"), p.get("dia"), p.get("cuotas_total") or "",
           p.get("inicio") or "", p.get("desde") or "",
           ", ".join(sorted(k for k, v in (p.get("pagados") or {}).items() if v)),
           "Sí" if p.get("ejemplo") else ""] for p in pagos],
         [24, 12, 6, 14, 18, 12, 40, 9])
    hoja(wb, "Trabajos",
         ["Fecha", "Cliente", "Maquinaria", "Cobro", "Cantidad", "Valor sin IVA", "Neto", "IVA", "Total", "Ejemplo"],
         [[t.get("fecha"), t.get("cliente"), t.get("maquinaria"),
           {"hora": "Por hora", "dia": "Por día", "trato": "Por trato"}.get(t.get("tipo"), t.get("tipo")),
           t.get("cantidad"), t.get("valor"), neto(t),
           round(neto(t) * IVA) if t.get("con_iva", True) else 0,
           round(neto(t) * (1 + IVA)) if t.get("con_iva", True) else neto(t),
           "Sí" if t.get("ejemplo") else ""] for t in trabajos],
         [12, 26, 20, 11, 9, 13, 12, 11, 12, 9])
    hoja(wb, "Gastos",
         ["Fecha", "En qué", "Monto", "Foto (carpeta Boletas)", "Ejemplo"],
         [[g.get("fecha"), g.get("descripcion"), g.get("monto"),
           ("boleta-" + g["foto"] + ".jpg") if g.get("foto") else "",
           "Sí" if g.get("ejemplo") else ""] for g in gastos],
         [12, 30, 12, 44, 9])
    wb.save(salida)
    print(f"{len(pagos)} pagos, {len(trabajos)} trabajos, {len(gastos)} gastos -> {salida}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
